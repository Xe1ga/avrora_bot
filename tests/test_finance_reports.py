"""Тесты XLSX-отчётов финансового учёта (форма вкладок ``tmp/avrora.xlsx``)."""

from collections.abc import Callable
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from avrora_bot.adapters.database.seed import seed_reference_data
from avrora_bot.adapters.reports.balances_xlsx import build_balances_report
from avrora_bot.adapters.reports.expenses_xlsx import build_expenses_report
from avrora_bot.adapters.reports.finance_xlsx import build_finance_workbook
from avrora_bot.adapters.reports.subscription_xlsx import (
    build_subscription_report,
)
from avrora_bot.application.use_cases.balances import (
    BalanceReport,
    BalanceRow,
)
from avrora_bot.application.use_cases.finance import ExpenseRow
from avrora_bot.application.use_cases.registration import (
    RegistrationData,
    RegistrationUseCases,
)
from avrora_bot.application.use_cases.roles import RoleUseCases
from avrora_bot.application.use_cases.subscriptions import (
    MonthSummary,
    PaymentRow,
    SubscriptionUseCases,
)
from avrora_bot.domain.entities import Expense, Subscription, User
from avrora_bot.domain.enums import ExpenseCategory, PaymentStatus, RoleName
from avrora_bot.domain.errors import PermissionDeniedError, ValidationError
from avrora_bot.domain.ports.uow import UnitOfWork
from avrora_bot.domain.services.balance_calc import AccountBalance
from avrora_bot.domain.value_objects import MonthPeriod
from tests.conftest import FakeVkGateway

SEPTEMBER = MonthPeriod(2026, 9)
GREEN = 'FF00FF00'
RED = 'FFFF0000'


def _payment(  # noqa: PLR0913
    name: str,
    *,
    paid: bool = True,
    amount: str = '2450',
    collector: str | None = 'Ольга Седова',
    note: str | None = None,
    day: int = 2,
) -> PaymentRow:
    return PaymentRow(
        user=User(vk_id=None, full_name=name),
        status=PaymentStatus.PAID if paid else PaymentStatus.UNPAID,
        amount=Decimal(amount),
        collector_name=collector if paid else None,
        note=note,
        marked_at=datetime(2026, 9, day, tzinfo=UTC) if paid else None,
    )


def _summary(rows: list[PaymentRow]) -> MonthSummary:
    return MonthSummary(
        subscription=Subscription(
            period_year=2026,
            period_month=9,
            total_amount=Decimal('65750'),
            voters_count=len(rows),
            per_person_amount=Decimal('2436'),
            per_percent_amount_fact=Decimal('2450'),
        ),
        rows=rows,
    )


def _sheet(path: Path):
    return load_workbook(path).worksheets[0]


def test_subscription_sheet_matches_template(tmp_path: Path) -> None:
    rows = [
        _payment('Света Яскевич', day=3),
        _payment('Оля Седова', day=1),
        _payment(
            'Алена Ворошилова',
            amount='2100',
            note='350 у Оли Тепловой (абонемент за вычетом разовой)',
            day=5,
        ),
        _payment('Ира Соколова', collector='Ольга Теплова', day=4),
        _payment('Анна Резухина', paid=False, note='обещала в пятницу'),
    ]
    ws = _sheet(
        build_subscription_report(
            SEPTEMBER, _summary(rows), tmp_path / 'sub.xlsx'
        )
    )

    assert ws.title == 'Сентябрь абонемент'
    assert ws['A1'].value == 'Сентябрь 2026'
    assert ws['A1'].font.sz == 12
    assert ws['A2'].value == 'Абонемент'
    assert {'A1:D1', 'A2:E2', 'A3:B3', 'A4:B4', 'A5:B5'} <= {
        str(r) for r in ws.merged_cells.ranges
    }
    # Два казначея — строки 4 и 5, шапка таблицы на 6-й.
    assert ws['A3'].value == 'ИТОГО:'
    assert ws['C3'].value == '=SUM(C7:C11)'
    assert (ws['D3'].value, ws['E3'].value) == ('ЦЕЛЬ:', 65750)
    assert (ws['D4'].value, ws['E4'].value) == ('НЕДОБОР:', '=E3-C3')
    assert ws['A4'].value == 'На счету Ольга Седова:'
    assert ws['C4'].value == '=SUMIF(D7:D11, "Ольга Седова", C7:C11)'
    assert ws['A5'].value == 'На счету Ольга Теплова:'
    assert [c.value for c in ws[6]][:5] == [
        '№',
        'ФИО',
        'Сумма',
        'Казначей (ФИО кто собирал)',
        'Примечание',
    ]
    # Оплатившие — в порядке оплаты, неоплатившие — в конце без суммы.
    table = [[c.value for c in ws[r]][:5] for r in range(7, 12)]
    assert table == [
        [1, 'Оля Седова', 2450, 'Ольга Седова', None],
        [2, 'Света Яскевич', 2450, 'Ольга Седова', None],
        [3, 'Ира Соколова', 2450, 'Ольга Теплова', None],
        [
            4,
            'Алена Ворошилова',
            2100,
            'Ольга Седова',
            '350 у Оли Тепловой (абонемент за вычетом разовой)',
        ],
        [5, 'Анна Резухина', None, None, 'не оплачено; обещала в пятницу'],
    ]
    assert ws['A3'].fill.fgColor.rgb == GREEN
    assert ws['B6'].font.b and ws['B6'].font.i
    assert ws['B7'].border.left.style == 'thin'


def test_subscription_sheet_without_payments(tmp_path: Path) -> None:
    ws = _sheet(
        build_subscription_report(SEPTEMBER, _summary([]), tmp_path / 's.xlsx')
    )
    assert ws['C3'].value == 0
    assert ws['A5'].value == '№'


def _expense(
    day: int, description: str, amount: str, account: str
) -> ExpenseRow:
    return ExpenseRow(
        expense=Expense(
            spent_on=date(2026, 9, day),
            category=ExpenseCategory.COACH,
            description=description,
            amount=Decimal(amount),
            account_user_id=1,
            created_by_vk_id=1,
            id=day,
        ),
        account_name=account,
    )


def test_expenses_sheet_matches_template(tmp_path: Path) -> None:
    rows = [
        _expense(1, 'Оплата зала сентябрь', '38750', 'Ольга Седова'),
        _expense(5, 'Оплата тренер Елена', '3000', 'Ольга Седова'),
        _expense(8, 'Мячи', '1200.50', 'Маргарита Фефилатьева'),
    ]
    ws = _sheet(build_expenses_report(SEPTEMBER, rows, tmp_path / 'e.xlsx'))

    assert ws.title == 'РАСХОДЫ'
    assert ws['A1'].value == 'РАСХОДЫ — сентябрь 2026'
    assert ws['A2'].value == 'ИТОГО:'
    assert ws['C2'].value == '=SUM(C6:C8)'
    assert ws['A3'].value == 'Со счёта Маргарита Фефилатьева:'
    assert ws['C3'].value == '=SUMIF(D6:D8, "Маргарита Фефилатьева", C6:C8)'
    assert ws['A4'].value == 'Со счёта Ольга Седова:'
    assert [c.value for c in ws[5]] == [
        'Дата',
        'Тип расхода',
        'Сумма',
        'С какого счета',
        'Примечание',
    ]
    assert ws['A6'].value == datetime(2026, 9, 1)
    assert ws['A6'].number_format == 'dd.mm.yyyy'
    assert [c.value for c in ws[8]][1:4] == [
        'Мячи',
        1200.5,
        'Маргарита Фефилатьева',
    ]
    assert ws['A5'].fill.fgColor.rgb == RED


def test_balances_sheet_sums_every_row(tmp_path: Path) -> None:
    report = BalanceReport(
        start=date(2026, 9, 1),
        end=date(2026, 9, 21),
        rows=[
            BalanceRow(
                name='Ольга Седова',
                balance=AccountBalance(
                    user_id=1,
                    opening=Decimal('1800'),
                    subscriptions=Decimal('60900'),
                    expenses=Decimal('61250'),
                ),
            ),
            BalanceRow(
                name='Ольга Теплова',
                balance=AccountBalance(user_id=2, one_time=Decimal('8050')),
            ),
        ],
    )
    ws = _sheet(build_balances_report(report, tmp_path / 'b.xlsx'))

    assert ws.title == 'ОСТАТОК'
    assert ws['A1'].value == 'ИТОГО по состоянию на 21.09.2026:'
    # В ручной таблице SUM не захватил последнюю строку — здесь все строки.
    assert ws['B1'].value == '=SUM(B3:B4)'
    assert ws['C2'].value == 'На 01.09.2026'
    assert [c.value for c in ws[3]] == [
        'Ольга Седова',
        '=C3+D3+E3-F3+G3+H3',
        1800,
        60900,
        0,
        61250,
        0,
        0,
    ]
    assert ws['A4'].value == 'Ольга Теплова'
    assert ws['E4'].value == 8050


def test_finance_workbook_has_all_tabs(tmp_path: Path) -> None:
    balances = BalanceReport(
        start=date(2026, 9, 1), end=date(2026, 9, 30), rows=[]
    )
    path = build_finance_workbook(
        SEPTEMBER,
        _summary([_payment('Оля Седова')]),
        [],
        [],
        balances,
        tmp_path / 'f.xlsx',
    )
    assert load_workbook(path).sheetnames == [
        'Сентябрь абонемент',
        'Сентябрь разовая',
        'РАСХОДЫ',
        'ОСТАТОК',
    ]
    path = build_finance_workbook(
        SEPTEMBER, None, [], [], balances, tmp_path / 'g.xlsx'
    )
    assert load_workbook(path).sheetnames == [
        'Сентябрь разовая',
        'РАСХОДЫ',
        'ОСТАТОК',
    ]
    assert load_workbook(path)['ОСТАТОК']['B1'].value == 0


@pytest.mark.asyncio
async def test_payment_fact_amount_and_note(
    uow_factory: Callable[[], UnitOfWork],
) -> None:
    vk = FakeVkGateway(admins={1})
    async with uow_factory() as uow:
        await seed_reference_data(uow)
        await uow.commit()
    reg = RegistrationUseCases(uow_factory, vk)
    for vk_id, name in [(101, 'Седова Ольга'), (102, 'Ворошилова Алена')]:
        await reg.self_register(RegistrationData(vk_id=vk_id, full_name=name))
        await reg.approve(1, vk_id)
    await RoleUseCases(uow_factory, vk).assign_role(1, 101, RoleName.COLLECTOR)
    subs = SubscriptionUseCases(uow_factory, vk)
    await subs.register_voting(101, SEPTEMBER, ['101', '102'])
    await subs.set_fact_amount(101, SEPTEMBER, Decimal('2450'))
    await subs.mark_payment(101, SEPTEMBER, '101', paid=True)
    await subs.mark_payment(101, SEPTEMBER, 'Ворошилова', paid=True)

    row = await subs.set_payment_amount(
        101, SEPTEMBER, 'Ворошилова', Decimal('2100')
    )
    assert row.amount == Decimal('2100')
    row = await subs.set_payment_note(
        101, SEPTEMBER, 'Ворошилова', '350 у Оли Тепловой'
    )
    assert row.note == '350 у Оли Тепловой'

    summary = await subs.month_summary(SEPTEMBER)
    assert summary.collected == Decimal('4550')
    by_name = {r.user.full_name: r for r in summary.rows}
    assert by_name['Ворошилова Алена'].note == '350 у Оли Тепловой'
    assert by_name['Ворошилова Алена'].marked_at is not None

    row = await subs.set_payment_amount(101, SEPTEMBER, 'Ворошилова', None)
    assert row.amount == Decimal('2450')
    with pytest.raises(ValidationError):
        await subs.set_payment_amount(
            101, SEPTEMBER, 'Ворошилова', Decimal('-1')
        )
    with pytest.raises(PermissionDeniedError):
        await subs.set_payment_note(102, SEPTEMBER, 'Ворошилова', 'x')
