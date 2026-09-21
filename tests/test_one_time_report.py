"""Тесты XLSX-отчёта по разовым посещениям (форма из образца клуба)."""

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from avrora_bot.adapters.reports.one_time_xlsx import build_one_time_report
from avrora_bot.application.use_cases.one_time import VisitRow
from avrora_bot.domain.entities import OneTimePayment
from avrora_bot.domain.value_objects import MonthPeriod

_HEADER_FILL = 'FF00FFFF'
_ROW_FILL = 'FF00FF00'


def _row(
    day: int,
    full_name: str,
    collector: str | None,
    note: str | None = None,
    amount: str = '350',
) -> VisitRow:
    return VisitRow(
        visit=OneTimePayment(
            user_id=1,
            visit_date=date(2026, 9, day),
            amount=Decimal(amount),
            note=note,
            id=day,
        ),
        full_name=full_name,
        collector_name=collector,
    )


_ROWS = [
    _row(2, 'Юля Мельник', 'Ольга Теплова', 'https://vk.ru/lulika1709'),
    _row(3, 'Наташа Каргалова', 'Маргарита Фефилатьева'),
    _row(4, 'Аня Соколова', 'Ольга Теплова', 'Александра Викторовна З.'),
    _row(5, 'Гость без оплаты', None),
]


def _build(tmp_path: Path, rows: list[VisitRow]):
    out = build_one_time_report(
        MonthPeriod(2026, 9), rows, tmp_path / 'report.xlsx'
    )
    return load_workbook(out).worksheets[0]


def test_report_layout_matches_the_club_template(tmp_path: Path) -> None:
    ws = _build(tmp_path, _ROWS)

    assert ws.title == 'Сентябрь разовая'
    assert ws['A1'].value == 'Разовая оплата'
    assert sorted(str(r) for r in ws.merged_cells.ranges) == [
        'A1:E1',
        'A2:B2',
        'A3:B3',
        'A4:B4',
    ]
    # Шапка таблицы — сразу за строками итогов (их 1 + число казначеев).
    assert [ws.cell(row=5, column=c).value for c in range(1, 6)] == [
        'Дата тренировки',
        'ФИО',
        'Сумма',
        'Казначей (ФИО кто собирал)',
        'Примечание',
    ]
    assert [round(ws.column_dimensions[c].width) for c in 'ABCDE'] == [
        18,
        32,
        21,
        29,
        32,
    ]


def test_report_rows_carry_visit_data_including_the_note(
    tmp_path: Path,
) -> None:
    ws = _build(tmp_path, _ROWS)

    assert ws['A6'].value.date() == date(2026, 9, 2)
    assert ws['A6'].number_format == 'dd.mm.yyyy'
    assert ws['B6'].value == 'Юля Мельник'
    assert ws['C6'].value == 350
    assert ws['D6'].value == 'Ольга Теплова'
    assert ws['E6'].value == 'https://vk.ru/lulika1709'
    # Неоплаченное посещение: казначея и примечания нет, строка всё равно в
    # отчёте и попадает в ИТОГО.
    assert ws['D9'].value is None
    assert ws['E9'].value is None


def test_report_totals_are_formulas_over_the_data_range(
    tmp_path: Path,
) -> None:
    ws = _build(tmp_path, _ROWS)

    assert ws['A2'].value == 'ИТОГО:'
    assert ws['C2'].value == '=SUM(C6:C9)'
    # По строке на каждого казначея месяца, по алфавиту.
    assert ws['A3'].value == 'На счету Маргарита Фефилатьева:'
    assert ws['C3'].value == ('=SUMIF(D6:D9, "Маргарита Фефилатьева", C6:C9)')
    assert ws['A4'].value == 'На счету Ольга Теплова:'
    assert ws['C4'].value == '=SUMIF(D6:D9, "Ольга Теплова", C6:C9)'


def test_report_keeps_the_template_colors(tmp_path: Path) -> None:
    ws = _build(tmp_path, _ROWS)

    assert ws['A1'].fill.fgColor.rgb == _HEADER_FILL
    assert ws['C5'].fill.fgColor.rgb == _HEADER_FILL
    assert ws['A5'].font.b and ws['A5'].font.i
    assert ws['A5'].font.name == 'Arial'
    assert ws['B6'].fill.fgColor.rgb == _ROW_FILL
    # Заливка строки данных тянется до конца видимой области листа (A..Z).
    assert ws.cell(row=6, column=26).fill.fgColor.rgb == _ROW_FILL
    assert ws['B6'].border.left.style == 'thin'


def test_report_without_visits_has_no_data_rows(tmp_path: Path) -> None:
    ws = _build(tmp_path, [])

    # Ни одного казначея — остаётся только строка ИТОГО, шапка идёт третьей.
    assert ws['A3'].value == 'Дата тренировки'
    # Диапазона данных нет, поэтому вместо формулы — ноль.
    assert ws['C2'].value == 0


def test_report_escapes_quotes_in_collector_name(tmp_path: Path) -> None:
    ws = _build(tmp_path, [_row(2, 'Гость', 'Анна "Аня" Петрова')])

    assert ws['C3'].value == ('=SUMIF(D5:D5, "Анна ""Аня"" Петрова", C5:C5)')
