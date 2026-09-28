"""Общее оформление XLSX-отчётов в стиле ручной таблицы клуба (avrora.xlsx).

Все листы — Arial 10, тонкие рамки, шапка и итоги полужирным курсивом с
заливкой листа (голубая у «разовой», зелёная у «абонемента», красная у
«РАСХОДОВ» и «ОСТАТКА»). Итоги — формулы Excel (``SUM``/``SUMIF``), а не
посчитанные числа: файл продолжают вести руками, и суммы должны
пересчитываться сами.
"""

from decimal import Decimal

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Цвета — ARGB, как в образце.
CYAN_FILL = PatternFill('solid', fgColor='FF00FFFF')
GREEN_FILL = PatternFill('solid', fgColor='FF00FF00')
RED_FILL = PatternFill('solid', fgColor='FFFF0000')

HEADER_FONT = Font(name='Arial', size=10, bold=True, italic=True)
TITLE_FONT = Font(name='Arial', size=12, bold=True, italic=True)
ROW_FONT = Font(name='Arial', size=10)
DATE_FORMAT = 'dd.mm.yyyy'

_THIN = Side(style='thin')
_NO_SIDE = Side()

# Excel ограничивает имя листа 31 символом.
MAX_SHEET_TITLE = 31


def style_cell(
    cell: Cell | MergedCell,
    font: Font,
    fill: PatternFill | None,
    *,
    horizontal: str | None,
    edges: str,
) -> None:
    """Шрифт, заливка, выравнивание и рамка по сторонам из ``edges`` (lrtb)."""
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if horizontal is not None:
        cell.alignment = Alignment(horizontal=horizontal)
    cell.border = Border(
        left=_THIN if 'l' in edges else _NO_SIDE,
        right=_THIN if 'r' in edges else _NO_SIDE,
        top=_THIN if 't' in edges else _NO_SIDE,
        bottom=_THIN if 'b' in edges else _NO_SIDE,
    )


def merged_band(  # noqa: PLR0913
    ws: Worksheet,
    row: int,
    first_col: int,
    last_col: int,
    value: object,
    *,
    fill: PatternFill | None,
    horizontal: str,
    font: Font = HEADER_FONT,
    bordered: bool = True,
) -> None:
    """Объединённый диапазон одной строки со значением и оформлением.

    Рамку получает каждая ячейка диапазона (иначе Excel нарисует её только
    вокруг левой ячейки); заливку и выравнивание — левая: объединение
    показывает её стиль, а границы берёт с краёв диапазона.
    """
    if last_col > first_col:
        ws.merge_cells(
            start_row=row,
            start_column=first_col,
            end_row=row,
            end_column=last_col,
        )
    ws.cell(row=row, column=first_col).value = value
    for column in range(first_col, last_col + 1):
        edges = ''
        if bordered:
            edges = 'lrtb' if column == first_col else 'tb'
            if column == last_col:
                edges = 'rtb' if column != first_col else 'lrtb'
        is_first = column == first_col
        style_cell(
            ws.cell(row=row, column=column),
            font,
            fill if is_first else None,
            horizontal=horizontal if is_first else None,
            edges=edges,
        )


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def column(index: int) -> str:
    """Буква колонки по номеру (1 → A)."""
    return get_column_letter(index)


def sum_formula(col: str, first_row: int, last_row: int) -> str | int:
    """``=SUM(C6:C34)`` — или 0, если строк данных нет."""
    if last_row < first_row:
        return 0
    return f'=SUM({col}{first_row}:{col}{last_row})'


def sumif_formula(  # noqa: PLR0913
    key_col: str,
    key: str,
    sum_col: str,
    first_row: int,
    last_row: int,
) -> str | int:
    """``=SUMIF(D6:D34, "ФИО", C6:C34)`` — или 0 без строк данных."""
    if last_row < first_row:
        return 0
    # Кавычка внутри ФИО закрыла бы строковый литерал формулы — в Excel
    # она экранируется удвоением.
    escaped = key.replace('"', '""')
    return (
        f'=SUMIF({key_col}{first_row}:{key_col}{last_row}, "{escaped}", '
        f'{sum_col}{first_row}:{sum_col}{last_row})'
    )


def amount(value: Decimal) -> int | float:
    """Сумма числом: целые рубли — как в образце, без лишних «.00»."""
    return int(value) if value == value.to_integral_value() else float(value)


def sheet_title(title: str) -> str:
    return title[:MAX_SHEET_TITLE]
