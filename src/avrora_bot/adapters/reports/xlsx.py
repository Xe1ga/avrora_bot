"""Генерация XLSX-отчёта по месяцу через openpyxl (ТЗ 3.2 п.6).

Полный отчёт формируется как документ и отправляется вложением. Рендер
таблицы в изображение не используется — данные должны копироваться.
"""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from avrora_bot.application.use_cases.subscriptions import MonthSummary
from avrora_bot.domain.enums import PaymentStatus
from avrora_bot.domain.value_objects import MonthPeriod

_HEADER_FILL = PatternFill('solid', fgColor='4F81BD')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_TITLE_FONT = Font(bold=True, size=14)
_PAID_FILL = PatternFill('solid', fgColor='C6EFCE')
_UNPAID_FILL = PatternFill('solid', fgColor='FFC7CE')


def build_month_report(
    period: MonthPeriod,
    summary: MonthSummary,
    hall_sum: Decimal,
    coach_sum: Decimal,
    out_path: Path,
) -> Path:
    """Создаёт XLSX-файл отчёта и возвращает путь к нему."""
    wb = Workbook()
    ws = wb.active
    ws.title = str(period)

    _write_title(ws, period)
    _write_table(ws, summary)
    _write_totals(ws, summary, hall_sum, coach_sum)
    _adjust_columns(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _write_title(ws: Worksheet, period: MonthPeriod) -> None:
    cell = ws.cell(row=1, column=1, value=f'Отчёт за {period.label()}')
    cell.font = _TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)


def _write_table(ws: Worksheet, summary: MonthSummary) -> None:
    header_row = 3
    headers = ('№', 'Участник', 'Статус оплаты')
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center')

    for idx, row in enumerate(summary.rows, start=1):
        r = header_row + idx
        is_paid = row.status is PaymentStatus.PAID
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=row.user.full_name)
        status_cell = ws.cell(
            row=r,
            column=3,
            value='Оплачено' if is_paid else 'Не оплачено',
        )
        status_cell.fill = _PAID_FILL if is_paid else _UNPAID_FILL


def _write_totals(
    ws: Worksheet,
    summary: MonthSummary,
    hall_sum: Decimal,
    coach_sum: Decimal,
) -> None:
    start = 3 + len(summary.rows) + 2
    rows: tuple[tuple[str, object], ...] = (
        ('Сумма абонемента, ₽', summary.subscription.effective_amount),
        ('Проголосовало', summary.subscription.voters_count),
        ('Оплатили', summary.paid_count),
        ('Собрано, ₽', summary.collected),
        ('К переводу за зал, ₽', hall_sum),
        ('К переводу тренеру, ₽', coach_sum),
    )
    for offset, (label, value) in enumerate(rows):
        r = start + offset
        label_cell = ws.cell(row=r, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)


def _adjust_columns(ws: Worksheet) -> None:
    widths = {'A': 22, 'B': 32, 'C': 16}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
