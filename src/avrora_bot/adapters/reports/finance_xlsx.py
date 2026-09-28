"""Сводная книга за месяц — все вкладки ручной таблицы клуба в одном файле.

Порядок листов как в ``tmp/avrora.xlsx``: «<Месяц> абонемент», «<Месяц>
разовая», «РАСХОДЫ», «ОСТАТОК». Лист абонемента пропускается, если
подписки на месяц нет.
"""

from pathlib import Path

from openpyxl import Workbook

from avrora_bot.adapters.reports.balances_xlsx import write_balances_sheet
from avrora_bot.adapters.reports.expenses_xlsx import write_expenses_sheet
from avrora_bot.adapters.reports.one_time_xlsx import write_one_time_sheet
from avrora_bot.adapters.reports.subscription_xlsx import (
    write_subscription_sheet,
)
from avrora_bot.application.use_cases.balances import BalanceReport
from avrora_bot.application.use_cases.finance import ExpenseRow
from avrora_bot.application.use_cases.one_time import VisitRow
from avrora_bot.application.use_cases.subscriptions import MonthSummary
from avrora_bot.domain.value_objects import MonthPeriod


def build_finance_workbook(  # noqa: PLR0913
    period: MonthPeriod,
    summary: MonthSummary | None,
    visits: list[VisitRow],
    expenses: list[ExpenseRow],
    balances: BalanceReport,
    out_path: Path,
) -> Path:
    """Создаёт сводную книгу за месяц и возвращает путь к ней."""
    wb = Workbook()
    wb.remove(wb.active)
    if summary is not None:
        write_subscription_sheet(wb.create_sheet(), period, summary)
    write_one_time_sheet(wb.create_sheet(), period, visits)
    write_expenses_sheet(wb.create_sheet(), period, expenses)
    write_balances_sheet(wb.create_sheet(), balances)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
