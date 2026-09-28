"""XLSX-отчёт по расходам за месяц — вкладка «РАСХОДЫ» (``tmp/avrora.xlsx``).

Красная шапка, ``ИТОГО`` и по строке «Со счёта <человек>» (``SUMIF``) на
каждого, с чьих денег платили, затем таблица «Дата | Тип расхода | Сумма |
С какого счета | Примечание». В колонку «Тип расхода» идёт описание
расхода («Оплата тренер Елена») — так её заполняли в таблице.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from avrora_bot.adapters.reports import xlsx_style as st
from avrora_bot.application.use_cases.finance import ExpenseRow
from avrora_bot.domain.value_objects import MonthPeriod

SHEET_TITLE = 'РАСХОДЫ'
_HEADERS = ('Дата', 'Тип расхода', 'Сумма', 'С какого счета', 'Примечание')
_COLUMN_WIDTHS = {'A': 19.63, 'B': 32.0, 'C': 23.13, 'D': 30.0, 'E': 40.0}
_LAST_COL = 5
_TITLE_ROW = 1
_TOTAL_ROW = 2


def build_expenses_report(
    period: MonthPeriod, rows: list[ExpenseRow], out_path: Path
) -> Path:
    """Создаёт XLSX-отчёт по расходам и возвращает путь к нему."""
    wb = Workbook()
    write_expenses_sheet(wb.active, period, rows)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def write_expenses_sheet(
    ws: Worksheet, period: MonthPeriod, rows: list[ExpenseRow]
) -> None:
    """Заполняет лист «РАСХОДЫ» за месяц ``period``."""
    ws.title = SHEET_TITLE
    accounts = sorted({row.account_name for row in rows})
    header_row = _TOTAL_ROW + len(accounts) + 1
    first_data_row = header_row + 1
    last_data_row = header_row + len(rows)

    st.merged_band(
        ws,
        _TITLE_ROW,
        1,
        _LAST_COL,
        f'РАСХОДЫ — {period.label()}',
        fill=st.RED_FILL,
        horizontal='center',
        bordered=False,
    )
    _write_total(
        ws,
        _TOTAL_ROW,
        'ИТОГО:',
        st.sum_formula('C', first_data_row, last_data_row),
    )
    for offset, name in enumerate(accounts, start=1):
        _write_total(
            ws,
            _TOTAL_ROW + offset,
            f'Со счёта {name}:',
            st.sumif_formula('D', name, 'C', first_data_row, last_data_row),
        )

    for col, name in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col)
        cell.value = name
        st.style_cell(
            cell, st.HEADER_FONT, st.RED_FILL, horizontal='center', edges='lrtb'
        )

    for offset, row in enumerate(rows):
        r = first_data_row + offset
        values = (
            row.expense.spent_on,
            row.expense.description,
            st.amount(row.expense.amount),
            row.account_name,
            row.expense.note,
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col)
            cell.value = value
            st.style_cell(
                cell, st.ROW_FONT, None, horizontal=None, edges='lrtb'
            )
        ws.cell(row=r, column=1).number_format = st.DATE_FORMAT

    st.set_widths(ws, _COLUMN_WIDTHS)


def _write_total(ws: Worksheet, row: int, label: str, value: object) -> None:
    st.merged_band(
        ws,
        row,
        1,
        2,
        label,
        fill=st.RED_FILL,
        horizontal='right',
        bordered=False,
    )
    cell = ws.cell(row=row, column=3)
    cell.value = value
    st.style_cell(
        cell, st.HEADER_FONT, st.RED_FILL, horizontal='center', edges=''
    )
