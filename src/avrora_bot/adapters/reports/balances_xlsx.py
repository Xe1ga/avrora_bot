"""XLSX-отчёт об остатках на дату — вкладка «ОСТАТОК» (``tmp/avrora.xlsx``).

Первые две колонки — как в образце: у кого и сколько на руках, над ними
«ИТОГО по состоянию на <дата>» (``SUM`` по всем строкам — в ручной
таблице диапазон забыл последнюю строку). Дальше — разбивка, из которой
остаток складывается формулой: на начало периода + абонементы + разовые
− расходы ± передачи ± корректировки.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from avrora_bot.adapters.reports import xlsx_style as st
from avrora_bot.application.use_cases.balances import BalanceReport

SHEET_TITLE = 'ОСТАТОК'
_COLUMN_WIDTHS = {
    'A': 32.0,
    'B': 16.0,
    'C': 16.0,
    'D': 16.0,
    'E': 16.0,
    'F': 16.0,
    'G': 16.0,
    'H': 18.0,
}
_TOTAL_ROW = 1
_HEADER_ROW = 2
_FIRST_DATA_ROW = 3


def build_balances_report(report: BalanceReport, out_path: Path) -> Path:
    """Создаёт XLSX-отчёт об остатках и возвращает путь к нему."""
    wb = Workbook()
    write_balances_sheet(wb.active, report)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def write_balances_sheet(ws: Worksheet, report: BalanceReport) -> None:
    """Заполняет лист «ОСТАТОК»."""
    ws.title = SHEET_TITLE
    last_data_row = _FIRST_DATA_ROW + len(report.rows) - 1

    label = ws.cell(row=_TOTAL_ROW, column=1)
    label.value = f'ИТОГО по состоянию на {report.end.strftime("%d.%m.%Y")}:'
    st.style_cell(
        label, st.HEADER_FONT, st.RED_FILL, horizontal=None, edges='lrtb'
    )
    total = ws.cell(row=_TOTAL_ROW, column=2)
    total.value = st.sum_formula('B', _FIRST_DATA_ROW, last_data_row)
    st.style_cell(
        total, st.HEADER_FONT, st.RED_FILL, horizontal='right', edges='lrtb'
    )

    headers = (
        'У кого',
        'Остаток',
        f'На {report.start.strftime("%d.%m.%Y")}',
        'Абонементы',
        'Разовые',
        'Расходы',
        'Передачи',
        'Корректировки',
    )
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=col)
        cell.value = name
        st.style_cell(
            cell, st.HEADER_FONT, st.RED_FILL, horizontal='center', edges='lrtb'
        )

    for offset, row in enumerate(report.rows):
        r = _FIRST_DATA_ROW + offset
        b = row.balance
        values = (
            row.name,
            f'=C{r}+D{r}+E{r}-F{r}+G{r}+H{r}',
            st.amount(b.opening),
            st.amount(b.subscriptions),
            st.amount(b.one_time),
            st.amount(b.expenses),
            st.amount(b.transfers),
            st.amount(b.adjustments),
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col)
            cell.value = value
            st.style_cell(
                cell,
                st.ROW_FONT,
                None,
                horizontal=None if col == 1 else 'right',
                edges='lrtb',
            )

    st.set_widths(ws, _COLUMN_WIDTHS)
