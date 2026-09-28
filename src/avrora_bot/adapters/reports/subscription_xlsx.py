"""XLSX-отчёт по абонементам за месяц — вкладка «<Месяц> абонемент».

Форма повторяет ручную таблицу клуба (``tmp/avrora.xlsx``): зелёная шапка
«Сентябрь 2026 / Абонемент», ``ИТОГО`` (сумма собранного), ``ЦЕЛЬ``
(стоимость месяца — зал + тренер), ``НЕДОБОР = ЦЕЛЬ − ИТОГО`` и таблица
«№ | ФИО | Сумма | Казначей (ФИО кто собирал) | Примечание». В отличие от
образца, под «ИТОГО» — строки «На счету <казначей>» (``SUMIF``), как на
вкладке разовых. Отменившиеся в отчёт не попадают — их убирают из
подписки («убрать голос»); не оплатившие пока идут в конце списка без
суммы, с пометкой в примечании.
"""

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from avrora_bot.adapters.reports import xlsx_style as st
from avrora_bot.application.use_cases.subscriptions import (
    MonthSummary,
    PaymentRow,
)
from avrora_bot.domain.enums import PaymentStatus
from avrora_bot.domain.value_objects import MonthPeriod

_HEADERS = (
    '№',
    'ФИО',
    'Сумма',
    'Казначей (ФИО кто собирал)',
    'Примечание',
)
_COLUMN_WIDTHS = {'A': 9.13, 'B': 32.0, 'C': 20.5, 'D': 35.63, 'E': 54.88}
_LAST_COL = 5
_UNPAID_NOTE = 'не оплачено'

_TITLE_ROW = 1
_SUBTITLE_ROW = 2
_TOTAL_ROW = 3
_SHORTFALL_ROW = 4


def build_subscription_report(
    period: MonthPeriod, summary: MonthSummary, out_path: Path
) -> Path:
    """Создаёт XLSX-отчёт по абонементам и возвращает путь к нему."""
    wb = Workbook()
    write_subscription_sheet(wb.active, period, summary)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def write_subscription_sheet(
    ws: Worksheet, period: MonthPeriod, summary: MonthSummary
) -> None:
    """Заполняет лист «<Месяц> абонемент»."""
    ws.title = st.sheet_title(f'{period.month_name.capitalize()} абонемент')
    rows = _ordered(summary.rows)
    collectors = sorted({r.collector_name for r in rows if _counted(r)})
    # Строки «На счету …» начинаются рядом с «НЕДОБОР» (строка 4); если
    # казначеев нет, строка 4 всё равно есть — как в образце.
    header_row = _SHORTFALL_ROW + max(len(collectors), 1)
    first_data_row = header_row + 1
    last_data_row = header_row + len(rows)

    st.merged_band(
        ws,
        _TITLE_ROW,
        1,
        4,
        f'{period.month_name.capitalize()} {period.year}',
        fill=st.GREEN_FILL,
        horizontal='center',
        font=st.TITLE_FONT,
        bordered=False,
    )
    st.merged_band(
        ws,
        _SUBTITLE_ROW,
        1,
        _LAST_COL,
        'Абонемент',
        fill=st.GREEN_FILL,
        horizontal='center',
    )
    _write_totals(ws, summary, collectors, first_data_row, last_data_row)
    _write_headers(ws, header_row)
    _write_rows(ws, rows, first_data_row)
    st.set_widths(ws, _COLUMN_WIDTHS)


def _counted(row: PaymentRow) -> bool:
    """Строка с суммой: оплачено и известно, кто собрал."""
    return row.status is PaymentStatus.PAID and row.collector_name is not None


def _ordered(rows: list[PaymentRow]) -> list[PaymentRow]:
    """Оплатившие — в порядке оплаты (как вносят в таблицу), затем прочие."""
    paid = [r for r in rows if r.status is PaymentStatus.PAID]
    unpaid = [r for r in rows if r.status is not PaymentStatus.PAID]
    paid.sort(key=lambda r: (_moment(r), r.user.full_name))
    unpaid.sort(key=lambda r: r.user.full_name)
    return paid + unpaid


def _moment(row: PaymentRow) -> datetime:
    moment = row.marked_at or datetime.min
    return moment if moment.tzinfo else moment.replace(tzinfo=UTC)


def _label(ws: Worksheet, row: int, col: int, value: object) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = value
    st.style_cell(
        cell, st.HEADER_FONT, st.GREEN_FILL, horizontal='center', edges='lrtb'
    )


def _write_totals(
    ws: Worksheet,
    summary: MonthSummary,
    collectors: list[str],
    first_data_row: int,
    last_data_row: int,
) -> None:
    st.merged_band(
        ws, _TOTAL_ROW, 1, 2, 'ИТОГО:', fill=st.GREEN_FILL, horizontal='center'
    )
    _label(
        ws, _TOTAL_ROW, 3, st.sum_formula('C', first_data_row, last_data_row)
    )
    _label(ws, _TOTAL_ROW, 4, 'ЦЕЛЬ:')
    _label(ws, _TOTAL_ROW, 5, st.amount(summary.subscription.total_amount))
    _label(ws, _SHORTFALL_ROW, 4, 'НЕДОБОР:')
    _label(ws, _SHORTFALL_ROW, 5, f'=E{_TOTAL_ROW}-C{_TOTAL_ROW}')

    if not collectors:
        st.merged_band(
            ws,
            _SHORTFALL_ROW,
            1,
            2,
            None,
            fill=st.GREEN_FILL,
            horizontal='right',
        )
    for offset, name in enumerate(collectors):
        row = _SHORTFALL_ROW + offset
        st.merged_band(
            ws,
            row,
            1,
            2,
            f'На счету {name}:',
            fill=st.GREEN_FILL,
            horizontal='right',
        )
        _label(
            ws,
            row,
            3,
            st.sumif_formula('D', name, 'C', first_data_row, last_data_row),
        )


def _write_headers(ws: Worksheet, header_row: int) -> None:
    for col, name in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col)
        cell.value = name
        st.style_cell(
            cell, st.HEADER_FONT, None, horizontal='center', edges='lrtb'
        )


def _write_rows(
    ws: Worksheet, rows: list[PaymentRow], first_data_row: int
) -> None:
    for offset, row in enumerate(rows):
        r = first_data_row + offset
        paid = row.status is PaymentStatus.PAID
        note = row.note
        if not paid:
            note = f'{_UNPAID_NOTE}; {note}' if note else _UNPAID_NOTE
        values = (
            offset + 1,
            row.user.full_name,
            st.amount(row.amount) if paid else None,
            row.collector_name if paid else None,
            note,
        )
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col)
            cell.value = value
            st.style_cell(
                cell, st.ROW_FONT, None, horizontal='left', edges='lrtb'
            )
